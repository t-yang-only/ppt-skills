"""Structural/text checks for the scientific visual baseline."""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation


def iter_text(slide):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                text = run.text.strip()
                if text:
                    yield text, run.font.size.pt if run.font.size else None


def is_metadata(text: str) -> bool:
    lowered = text.lower()
    return any(token in lowered for token in ("doi", "paper 03", "source:", "nature 634", "conceptual cover")) or bool(
        re.fullmatch(r"\d+\s*/\s*\d+", text)
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--expected-slides", type=int, default=12)
    args = parser.parse_args()
    if not args.pptx.exists():
        print(f"FAIL: missing PPTX: {args.pptx}")
        return 2

    prs = Presentation(str(args.pptx))
    width = round(prs.slide_width / 914400, 3)
    height = round(prs.slide_height / 914400, 3)
    failures = []
    warnings = []
    if len(prs.slides) != args.expected_slides:
        failures.append(f"slide count {len(prs.slides)} != expected {args.expected_slides}")
    if abs(width / height - 16 / 9) > 0.01:
        failures.append(f"page size {width} × {height} is not 16:9")

    page_numbers = set()
    for index, slide in enumerate(prs.slides, start=1):
        texts = list(iter_text(slide))
        all_text = "\n".join(text for text, _ in texts)
        if re.search(rf"\b{index}\s*/\s*{len(prs.slides)}\b", all_text):
            page_numbers.add(index)
        elif index > 1:
            warnings.append(f"slide {index}: page number text not detected")
        for text, size in texts:
            if size is not None and size < 12 and not is_metadata(text):
                warnings.append(f"slide {index}: text below 12 pt ({size:g} pt): {text[:60]}")
        if index > 1 and not re.search(r"(?:source:|fig\.)", all_text, re.IGNORECASE):
            warnings.append(f"slide {index}: source/figure attribution not detected")

    expected_page_numbers = set(range(2, len(prs.slides) + 1))
    if page_numbers and page_numbers != expected_page_numbers:
        failures.append(f"page numbers are not continuous for content slides: {sorted(page_numbers)}")

    print(f"PPTX: {args.pptx}")
    print(f"PASS: {len(prs.slides)} slides, {width} × {height} in")
    print(f"WARNINGS: {len(warnings)}")
    for warning in warnings:
        print(f"  - {warning}")
    print(f"FAILURES: {len(failures)}")
    for failure in failures:
        print(f"  - {failure}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
