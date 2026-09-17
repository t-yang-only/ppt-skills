"""Build the neutral comparison baseline for the new High Line case."""
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "high_line_public_realm.pptx"
OUT = ROOT / "baseline" / "high_line_public_realm_baseline.pptx"

def add_text(slide, value, x, y, w, h, size=11, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.text = value or "[no text extracted]"
    for p in box.text_frame.paragraphs:
        p.font.name = "Arial"; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = RGBColor(25, 25, 25)

def main():
    source = Presentation(str(SOURCE)); baseline = Presentation()
    baseline.slide_width = source.slide_width; baseline.slide_height = source.slide_height
    blank = baseline.slide_layouts[6]
    for index, src in enumerate(source.slides, 1):
        slide = baseline.slides.add_slide(blank)
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor(248, 248, 246)
        texts = [shape.text.strip() for shape in src.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
        add_text(slide, f"BASELINE / HIGH LINE / {index:02d}", .55, .35, 12, .25, 8, True)
        add_text(slide, texts[0] if texts else f"Slide {index}", .65, .85, 5.9, .75, 22, True)
        add_text(slide, "\n".join(texts[1:]), .72, 1.8, 5.55, 4.7, 11)
        pictures = [shape for shape in src.shapes if getattr(shape, "shape_type", None) == 13]
        if pictures:
            slide.shapes.add_picture(BytesIO(pictures[0].image.blob), Inches(6.65), Inches(1.8), width=Inches(5.8), height=Inches(4.55))
        else:
            shape = slide.shapes.add_shape(1, Inches(6.65), Inches(1.8), Inches(5.8), Inches(4.55))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(225, 225, 220)
            shape.line.color.rgb = RGBColor(170, 170, 165)
            add_text(slide, "NEUTRAL VISUAL FIELD", 7.1, 3.85, 4.9, .4, 12, True)
        add_text(slide, "Independent neutral baseline · same content and page count", .65, 7.08, 12, .2, 7)
    OUT.parent.mkdir(parents=True, exist_ok=True); baseline.save(str(OUT)); print(OUT)

if __name__ == "__main__": main()
